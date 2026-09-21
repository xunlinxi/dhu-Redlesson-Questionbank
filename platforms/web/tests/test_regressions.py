import io
from pathlib import Path
import pytest
from docx import Document
from backend.parser import QuestionParser, parse_file


def test_word_tables_in_document_order(tmp_path):
    doc = Document()
    doc.add_paragraph('一、单项选择题')
    table = doc.add_table(rows=3, cols=1)
    for cell, text in zip([r.cells[0] for r in table.rows], ['1、表格题（A）', 'A.甲', 'B.乙']):
        cell.text = text
    path = tmp_path / 'table.docx'
    doc.save(path)
    q, *_ = parse_file(str(path), '表格')
    assert len(q) == 1
    assert q[0]['options'] == {'A': '甲', 'B': '乙'}


def test_standalone_separated_answer(tmp_path):
    path = tmp_path / 'test.txt'
    path.write_text('二、多项选择题\n1、选择\nA.甲\nB.乙\nC.丙\n答案：A、C', encoding='utf-8')
    q, *_ = parse_file(str(path), '测试')
    assert q[0]['answer'] == ['A', 'C']


def test_utf16_and_bom(tmp_path):
    for encoding in ['utf-16', 'utf-8-sig', 'gb18030']:
        path = tmp_path / 'test.txt'
        path.write_text('1、测试（A）\nA.甲\nB.乙', encoding=encoding)
        q, *_ = parse_file(str(path), '测试')
        assert len(q) == 1
        assert q[0]['answer'] == ['A']


def test_bank_ids_do_not_collide(tmp_path):
    path = tmp_path / 'test.txt'
    path.write_text('1、测试（A）\nA.甲\nB.乙', encoding='utf-8')
    a, *_ = parse_file(str(path), '自定义一')
    b, *_ = parse_file(str(path), '自定义二')
    assert a[0]['id'] != b[0]['id']


@pytest.mark.parametrize('query', ['count=-1', 'count=abc', 'single_count=-2', 'judge_count=x'])
def test_invalid_practice_counts(client, monkeypatch, query):
    from backend.models.questions import QuestionsModel
    monkeypatch.setattr(QuestionsModel, 'get_questions', lambda **kw: [{'id': 'a', 'type': 'single'}])
    response = client.get('/api/practice/random?' + query)
    assert response.status_code == 400
    assert response.json['success'] is False

@pytest.fixture
def isolated_import(client, monkeypatch, tmp_path):
    from backend.models.questions import QuestionsModel
    from backend.routes import banks
    monkeypatch.setattr(QuestionsModel, 'get_file_path', staticmethod(lambda: str(tmp_path / 'questions.json')))
    uploads = tmp_path / 'uploads'
    uploads.mkdir()
    monkeypatch.setattr(banks, 'UPLOAD_FOLDER', str(uploads))
    return client, uploads


def test_docx_table_upload_route(isolated_import):
    client, uploads = isolated_import
    document = Document()
    document.add_paragraph('一、单项选择题')
    table = document.add_table(rows=3, cols=1)
    for row, text in zip(table.rows, ['1、表格题（B）', 'A.甲', 'B.乙']):
        row.cells[0].text = text
    buffer = io.BytesIO()
    document.save(buffer)
    buffer.seek(0)
    response = client.post('/api/import', data={'file': (buffer, '表格.docx'), 'bank_name': '表格'})
    assert response.status_code == 200
    assert response.json['question_count'] == 1
    assert list(uploads.iterdir()) == []


def test_invalid_import_keeps_previous_bank(isolated_import):
    client, uploads = isolated_import
    for text, expected in [('1、测试（A）\nA.甲\nB.乙', 200), ('完全没有题目', 400)]:
        response = client.post('/api/import', data={'file': (io.BytesIO(text.encode()), 'test.txt'), 'bank_name': '保留'})
        assert response.status_code == expected
        assert list(uploads.iterdir()) == []
    questions = client.get('/api/questions?bank=保留').json['questions']
    assert len(questions) == 1
    assert questions[0]['answer'] == ['A']


def test_partial_import_reports_skipped_questions(isolated_import):
    client, _ = isolated_import
    text = '1、有效（A）\nA.甲\nB.乙\n2、无答案（）\nA.甲\nB.乙'
    response = client.post('/api/import', data={'file': (io.BytesIO(text.encode()), 'test.txt')})
    assert response.json['question_count'] == 1
    assert len(response.json['warnings']) == 1


@pytest.mark.parametrize('keyword,count', [('毛概',345), ('纲要',550), ('习近平',367), ('思想道德',278)])
def test_shipped_corpus(keyword, count):
    root = Path(__file__).resolve().parents[2] / 'files'
    path = next(p for p in root.glob('*.txt') if keyword in p.name)
    questions, *_ = parse_file(str(path))
    assert len(questions) == count
