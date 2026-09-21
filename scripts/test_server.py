"""Isolated local UI test server; does not read/write the user's study records."""
import sys, tempfile, json
from pathlib import Path
ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'platforms/web'))
# Recreate ignored binary fixture for the offline browser regression.
from docx import Document
artifacts = ROOT / 'test-environment/artifacts'
artifacts.mkdir(parents=True, exist_ok=True)
document = Document()
document.add_paragraph('一、单项选择题')
table = document.add_table(rows=3, cols=1)
for row, text in zip(table.rows, ['1、表格里的题目（A）', 'A.正确选项', 'B.其他选项']):
    row.cells[0].text = text
document.save(artifacts / 'import-table.docx')

from backend import config
from backend.app import app
with tempfile.TemporaryDirectory(prefix='quiz-ui-') as directory:
    config.CONFIG_FILE=str(Path(directory)/'config.json')
    Path(config.CONFIG_FILE).write_text(json.dumps({'data_path':directory}),encoding='utf8')
    from flask import send_from_directory
    @app.route('/static-test/<path:name>')
    def static_test(name):
        return send_from_directory(ROOT/'platforms/web/static-site',name)
    @app.route('/android-test/<path:name>')
    def android_test(name):
        return send_from_directory(ROOT/'platforms/android/frontend',name)
    app.run(port=50123,host='127.0.0.1',threaded=True,use_reloader=False)
